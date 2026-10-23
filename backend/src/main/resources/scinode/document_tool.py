"""SciNode document nodes. Inputs remain intact; generated documents live in run output."""
import json, os, hashlib, zipfile
from pathlib import Path
from datetime import datetime, timezone

params = json.loads(Path(os.environ['SCINODE_INPUT']).read_text())
# Resolve recorded paths at runtime while preserving immutable historical manifests.
def relocated(value):
    old, new = os.environ.get('SCINODE_PREVIOUS_ROOT', ''), os.environ.get('SCINODE_CURRENT_ROOT', '')
    if isinstance(value, dict): return {k: relocated(v) for k, v in value.items()}
    if isinstance(value, list): return [relocated(v) for v in value]
    if isinstance(value, str) and old and new and (value == old or value.startswith(old + '/')):
        return new + value[len(old):]
    return value
params = relocated(params)

out = Path(os.environ['SCINODE_OUTPUT'])
out.mkdir(exist_ok=True)
op = params['operation']
if op == 'document':
    path = Path(params['filePath'])
    suffix = path.suffix.lower()
    if suffix == '.docx':
        from docx import Document
        doc = Document(path)
        paragraphs = [{'index': i, 'text': p.text} for i, p in enumerate(doc.paragraphs)]
        tables = [[[c.text for c in row.cells] for row in t.rows] for t in doc.tables]
        text = '\n'.join(p['text'] for p in paragraphs)
        text += '\n' + '\n'.join('\t'.join(row) for table in tables for row in table)
        result = {'text': text, 'paragraphs': paragraphs, 'tables': tables}
    elif suffix == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(path)
        pages = [page.extract_text() or '' for page in reader.pages]
        if not any(s.strip() for s in pages):
            raise ValueError('PDF缺少可提取文本层，请先完成文字识别')
        result = {'text': '\n'.join(pages), 'pageCount': len(pages)}
    elif suffix in ('.txt', '.md', '.markdown', '.csv', '.json'):
        result = {'text': path.read_text(encoding='utf-8')}
    else:
        raise ValueError('文档格式请使用docx、pdf、txt、md、csv或json')
    result.update({'sourceName': path.name, 'sourceSha256': hashlib.sha256(path.read_bytes()).hexdigest()})
    (out / 'extracted.json').write_text(json.dumps(result, ensure_ascii=False), encoding='utf-8')
elif op == 'export':
    text = params.get('text', '')
    if not isinstance(text, str):
        raise ValueError('导出text输入需要文字类型，请选择上游text字段')
    if not text.strip():
        raise ValueError('请映射待导出的text字段')
    fmt = params.get('format') or 'docx'
    if fmt == 'docx':
        from docx import Document
        # Optional indexed patches retain existing figures/tables; text-only export creates a new report.
        if params.get('filePath') and params.get('patches'):
            doc = Document(params['filePath'])
            patches = params['patches']
            if isinstance(patches, str):
                raw = patches.strip()
                if raw.startswith('```'):
                    raw = raw.split('\n', 1)[1].rsplit('```', 1)[0]
                patches = json.loads(raw)
            patches = patches.get('patches', patches) if isinstance(patches, dict) else patches
            seen = set()
            for patch in patches:
                if 'original' not in patch or not isinstance(patch.get('text'), str):
                    raise ValueError('每项修订需提供index、original与text')
                idx = int(patch['index'])
                if idx in seen or idx < 0 or idx >= len(doc.paragraphs):
                    raise ValueError(f'段落索引重复或越界: {idx}')
                seen.add(idx)
                p = doc.paragraphs[idx]
                if p.text != patch['original']:
                    raise ValueError(f'段落原文校验失败: {idx}')
                if p._p.xpath('.//w:drawing'):
                    raise ValueError(f'含图段落请单独处理: {idx}')
                if p.runs:
                    p.runs[0].text = patch['text']
                    for run in p.runs[1:]: run.text = ''
                else: p.add_run(patch['text'])
        else:
            doc = Document()
            for line in text.splitlines(): doc.add_paragraph(line)
        path = out / 'report.docx'
        doc.save(path)
    elif fmt == 'pdf':
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from xml.sax.saxutils import escape
        pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
        style = getSampleStyleSheet()['BodyText']
        style.fontName = 'STSong-Light'; style.leading = 18
        path = out / 'report.pdf'
        SimpleDocTemplate(str(path)).build([Paragraph(escape(s) or ' ', style) for s in text.splitlines()])
    elif fmt in ('md', 'txt'):
        path = out / ('report.' + fmt)
        path.write_text(text, encoding='utf-8')
    else: raise ValueError('导出格式请使用docx、pdf、md或txt')
    result = {'text': '已导出 ' + path.name, 'fileName': path.name, 'exportMode': 'indexed_patches' if params.get('patches') else 'new_report'}
elif op == 'package':
    path = out / 'research-package.zip'
    manifest = {'createdAt': datetime.now(timezone.utc).isoformat(), 'files': []}
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('results.json', json.dumps(params, ensure_ascii=False, indent=2))
        for i, artifact in enumerate(params.get('artifacts', [])):
            source = Path(artifact['path'])
            actual = hashlib.sha256(source.read_bytes()).hexdigest()
            if actual != artifact['sha256']:
                raise ValueError('成果文件校验失败: ' + source.name)
            name = f'files/{i:03d}_{source.name}'
            archive.write(source, name)
            manifest['files'].append({'name': name, 'sha256': actual, 'size': source.stat().st_size})
        archive.writestr('manifest.json', json.dumps(manifest, ensure_ascii=False, indent=2))
    result = {'text': '成果包已生成', 'fileName': path.name, 'manifest': manifest}
else: raise ValueError('未知文档操作: ' + op)
(out / 'result.json').write_text(json.dumps(result, ensure_ascii=False), encoding='utf-8')
print(json.dumps(result, ensure_ascii=False))
